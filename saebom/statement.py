# [파일 기능명세]
# 계약서 → 거래명세서 엑셀 생성 엔진. contract-statement 스킬(extract_contract.py·fill_statement.py)의 규칙을 앱에 맞게 옮겼다.
# 차이: PDF를 다시 읽지 않고 앱이 저장한 문서의 extraction.pages(pdfplumber 페이지 텍스트)와 extraction.fields를 쓰며,
#       PDF 변환(LibreOffice)은 하지 않고, CLI가 아니라 함수로 제공한다. OCR·AI·네트워크를 사용하지 않는다.
# 처리 흐름: draft(문서, 계약 행) → 확인 창 초기값(수입 여부·거래처·작성 일자·청구 범위·품목·안내) → 사람이 화면에서 고침
#           → build(값, 저장 폴더) → 검사 통과 시 템플릿 사본을 채워 새 파일로 저장(덮어쓰기 금지, 원본 템플릿 불변).
# 템플릿 칸: A3 "{거래처}"·A5 "{YYYY}년 {MM}월 {DD}일" 자리표시자, 13~24행(12줄)의 A~E 열만 쓴다.
#           F~H 열과 25행 합계 수식, E3:F8·H2 공급자 칸은 건드리지 않으며 저장 전에 바뀌지 않았는지 검사한다.
# 유지보수: 템플릿 양식·규칙이 바뀌면 이 명세, tests/test_statement.py, 기획/거래명세서_생성_구축계획.md를 함께 갱신한다.

import calendar
import re
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from .compare import BIZ, DATE, integer
from . import config
from .files import safe_name
from .rules import MAX_AMOUNT, is_amount, is_month, spreadsheet_safe  # noqa: F401  MAX_AMOUNT는 statement.MAX_AMOUNT 참조 호환용으로 다시 노출

# 템플릿 파일 이름과 채우는 행 범위(13~24행 12줄, 25행은 총계 수식).
TEMPLATE_NAME = "거래명세서_템플릿.xlsx"
FIRST_ROW, LAST_ROW, TOTAL_ROW = 13, 24, 25
CAPACITY = LAST_ROW - FIRST_ROW + 1
# 월 금액 상한(원)은 rules.MAX_AMOUNT(공용 값)를 쓰며, 위 import로 statement.MAX_AMOUNT 이름을 유지한다.
# 수입 계약이 아닐 때 확인 창과 생성 거절에 쓰는 안내 문구.
SPEND_REASON = "지출 계약은 거래처가 발행하는 문서라 거래명세서를 만들지 않습니다."
UNKNOWN_REASON = "계약서 서명란에서 새봄바이오가 \"을\"(공급자)인지 확인하지 못해 거래명세서를 만들 수 없습니다."
NO_SOURCE = "계약서 원문 PDF가 없어 거래명세서를 만들 수 없습니다."
# 표지·본문의 따옴표는 곧은 따옴표와 둥근 따옴표가 섞여 있다.
QUOTE = r"[\"“”]"


# [기능명세] template_path
# 목적: 거래명세서 템플릿 파일 경로 계산
# 입력: ()
# 출력: Path — <루트>/saebom/assets/거래명세서_템플릿.xlsx
# 처리: PyInstaller 실행이면 sys._MEIPASS(번들 압축 해제 폴더)를, 소스 실행이면 saebom 패키지의 상위 폴더를 루트로 쓴다
#       (PyInstaller 빌드는 --add-data "saebom/assets;saebom/assets"로 같은 상대 경로에 넣는다)
# 예외/실패: 존재 여부는 확인하지 않음(없으면 build에서 openpyxl 오류)
# 부수효과: 없음
def template_path():
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))
    return root / "saebom" / "assets" / TEMPLATE_NAME


# [기능명세] page_text
# 목적: 저장된 문서의 페이지 텍스트를 하나의 원문 문자열로 결합
# 입력: (doc: Store 문서 사전 — extraction.pages는 [{page, text}] 또는 문자열 목록)
# 출력: 페이지 텍스트를 줄바꿈으로 이은 문자열(페이지가 없으면 빈 문자열)
# 처리: 사전이면 text 값을, 문자열이면 그대로 사용
# 예외/실패: 없음(형식이 다른 항목은 문자열로 변환)
# 부수효과: 없음
def page_text(doc):
    pages = (doc.get("extraction") or {}).get("pages") or []
    return "\n".join((p.get("text") or "") if isinstance(p, dict) else str(p) for p in pages)


# [기능명세] field
# 목적: 문서 추출값(extraction.fields)의 값 하나 조회
# 입력: (doc: 문서 사전, key: 필드 이름)
# 출력: 값 또는 None
# 처리: extraction → fields → key → value 순서로 안전 조회
# 예외/실패: 없음(누락은 None)
# 부수효과: 없음
# 통일하지 않은 이유: compare.val은 fields나 key 값이 None이면 AttributeError를 내지만 이 함수는 모든 단계의 None·빈 값을 견디고(문서 레코드를 받아 extraction을 한 겹 벗김), 의미가 달라 그대로 둔다.
def field(doc, key):
    return (((doc.get("extraction") or {}).get("fields") or {}).get(key) or {}).get("value")


# [기능명세] month_of
# 목적: ISO 날짜 문자열에서 청구 월(YYYY-MM) 추출
# 입력: (value: "YYYY-MM-DD" 또는 None)
# 출력: "YYYY-MM" 또는 None
# 처리: date.fromisoformat으로 검증한 뒤 앞 7자를 사용
# 예외/실패: 형식이 틀리면 None
# 부수효과: 없음
def month_of(value):
    try:
        return date.fromisoformat(value).isoformat()[:7] if value else None
    except (TypeError, ValueError):
        return None


# [기능명세] parse_month
# 목적: 화면이 보낸 청구 월 문자열 검증
# 입력: (value: "YYYY-MM", label: 오류 문구에 쓸 항목 이름)
# 출력: (연, 월) 정수 튜플
# 처리: 네 자리 연도-두 자리 월 형식 확인 후 rules.is_month로 실제 달력 월(1~12월, 0000년 제외) 확인
# 예외/실패: 형식·범위 오류는 ValueError(한글 안내)
# 부수효과: 없음
def parse_month(value, label):
    m = re.fullmatch(r"(\d{4})-(\d{2})", value or "")
    if not m or not is_month(value):
        raise ValueError(f"{label}을 YYYY-MM 형식으로 입력해 주세요.")
    return int(m.group(1)), int(m.group(2))


# [기능명세] month_range
# 목적: 시작 월부터 끝 월까지의 청구 월 목록 생성
# 입력: (lo: (연, 월), hi: (연, 월))
# 출력: [(연, 월), ...] — lo > hi이면 빈 목록
# 처리: 12월 다음은 다음 해 1월로 넘김
# 예외/실패: 없음
# 부수효과: 없음
def month_range(lo, hi):
    out, (y, m) = [], lo
    while (y, m) <= hi:
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


# [기능명세] safe_cell
# 목적: 엑셀 문자열 칸의 수식 해석 방지(CSV와 같은 규칙을 쓰도록 rules.spreadsheet_safe에 위임)
# 입력: (value: str)
# 출력: =, +, -, @, 탭, 캐리지 리턴으로 시작하면 앞에 작은따옴표를 붙인 문자열, 아니면 그대로
# 처리: openpyxl은 =로 시작하는 문자열을 수식으로 저장하므로 접두어로 막는다(탭·CR 접두까지 막는 것은 CSV와 통일하면서 추가된 의도된 변경)
# 예외/실패: 없음
# 부수효과: 없음
def safe_cell(value):
    return spreadsheet_safe(value)


# [기능명세] find_customer
# 목적: 계약서 원문에서 거래처(갑) 상호 찾기
# 입력: (text: 계약서 원문, supplier: 표지의 을 상호 또는 None)
# 출력: 갑 상호 문자열 또는 ""
# 처리: 1) 표지 "갑" : 상호 → 2) 서명란 "상 호 <갑> <을>" 줄에서 을 상호를 뒤에서 떼어 냄 → 3) 본문 "<상호>(이하 "갑"" 순서로 찾는다
# 예외/실패: 모두 없으면 빈 문자열(호출자가 missing에 기록)
# 부수효과: 없음
def find_customer(text, supplier):
    m = re.search(QUOTE + r"갑" + QUOTE + r"\s*:\s*([^\n]+)", text)
    if m and m.group(1).strip():
        return m.group(1).strip()
    m = re.search(r"상\s*호\s+([^\n]+)", text)
    if m and supplier:
        line = m.group(1).strip()
        if line.endswith(supplier) and line[:-len(supplier)].strip():
            return line[:-len(supplier)].strip()
    m = re.search(r"([^\n\"“”]+?)\s*\(이하\s*" + QUOTE + r"갑" + QUOTE, text)
    return m.group(1).strip() if m else ""


# [기능명세] draft
# 목적: 계약 행과 저장 문서로 거래명세서 확인 창의 초기값 생성
# 입력: (doc: Store 문서 사전(extraction.pages 원문 포함), row: contract_imports 계약 행 {id, kind, ...})
# 출력: {contract_id, eligible, reason, customer, customer_biz, contract_date, issue_date, period_start, period_end,
#        from_month, to_month, items:[{name, spec, monthly_amount}], notes:[], missing:[]}
# 처리: 서명란 "사업자등록번호 <갑> <을>"에서 두 번호를 읽어 을 번호 == config.OWN_BIZ이고 행 kind == "수입"이면 eligible
#       → 거래처는 표지 "갑" 상호(없으면 서명란), customer_biz는 갑 번호 → 계약체결일(compare.DATE)을 작성 일자 기본값으로
#       → 기간은 extraction.fields start/end와 그 월 → 품목: 이름이 붙은 "N. 이름: 월 금 X원" 줄이 2개 이상이면 줄마다 품목(규격 빈칸),
#         아니면 마지막 "별첨 1" 이후의 서비스명·대상과 fields.amount로 1개 → 총액 계약·부가세 포함 안내와 누락 항목 기록
# 예외/실패: 찾지 못한 값은 None/""로 두고 missing에 키를 남긴다(예: "contract_date", "items[0].spec"); 예외를 내지 않음
# 부수효과: 없음(입력 사전을 바꾸지 않음)
def draft(doc, row):
    text = page_text(doc)
    notes, missing = [], []
    sign = re.search(r"사업자등록번호\s+(" + BIZ + r")\s+(" + BIZ + r")", text)
    customer_biz, supplier_biz = (sign.group(1), sign.group(2)) if sign else (None, None)
    eligible = supplier_biz == config.OWN_BIZ and row.get("kind") == "수입"
    if eligible:
        reason = None
    elif row.get("kind") == "지출" or customer_biz == config.OWN_BIZ:
        reason = SPEND_REASON
    else:
        # 수입 계약인데 서명란을 읽지 못했거나 을 번호가 자사가 아닌 경우: 지출 안내 대신 확인 불가 안내를 준다.
        reason = UNKNOWN_REASON

    m = re.search(QUOTE + r"을" + QUOTE + r"\s*:\s*([^\n]+)", text)
    customer = find_customer(text, m.group(1).strip() if m else None)
    if not customer:
        missing.append("customer")

    contract_date = None
    m = re.search(r"계약체결일\s*:\s*" + DATE, text)
    if m:
        try:
            contract_date = date(*map(int, m.groups())).isoformat()
        except ValueError:
            contract_date = None
    if not contract_date:
        missing.append("contract_date")

    # 기간은 compare.extract_contract가 이미 ISO 날짜로 뽑아 둔 값을 재사용한다(형식이 틀리면 없는 것으로 본다).
    from_month, to_month = month_of(field(doc, "start")), month_of(field(doc, "end"))
    period_start = field(doc, "start") if from_month else None
    period_end = field(doc, "end") if to_month else None
    if not period_start:
        missing.append("period_start")
    if not period_end:
        missing.append("period_end")

    basis, amount = field(doc, "basis"), field(doc, "amount")
    lines = re.findall(r"(?m)^\s*\d+\.\s*([^\n:：]+?)\s*[:：]\s*월\s*금\s*(\d[\d,]*)\s*원", text)
    if len(lines) >= 2:
        items = [{"name": name.strip(), "spec": "", "monthly_amount": integer(value)} for name, value in lines]
        notes.append(f"서비스료 항목이 {len(items)}개입니다. 각 품목의 규격을 견적서에서 확인해 입력하세요.")
    else:
        # 본문 목록("* 별첨 1. 서비스의 범위")이 아니라 실제 별첨1 표를 보도록 마지막 등장 이후만 검색한다.
        annex = text.rsplit("별첨 1", 1)[-1] if "별첨 1" in text else ""
        name = re.search(r"서비스명\s+([^\n]+)", annex)
        spec = re.search(r"\d\s+대상\s+([^\n]+)", annex)
        items = [{"name": name.group(1).strip() if name else "", "spec": spec.group(1).strip() if spec else "",
                  "monthly_amount": amount if isinstance(amount, int) else None}]
    if basis == "total":
        notes.append("총액 계약입니다. 월금액을 직접 확인해 입력하세요.")
        for it in items:
            it["monthly_amount"] = None
    if field(doc, "vat_mode") == "included":
        notes.append("부가가치세 포함 금액입니다. 공급가액(부가세 제외)으로 고쳐 입력하세요.")
    for i, it in enumerate(items):
        for key in ("name", "spec", "monthly_amount"):
            if it[key] in (None, ""):
                missing.append(f"items[{i}].{key}")

    return {"contract_id": row.get("id"), "eligible": eligible, "reason": reason, "customer": customer, "customer_biz": customer_biz,
            "contract_date": contract_date, "issue_date": contract_date, "period_start": period_start, "period_end": period_end,
            "from_month": from_month, "to_month": to_month, "items": items, "notes": notes, "missing": missing}


# [기능명세] validate
# 목적: 생성 요청 값의 필수값·형식 검사와 정리
# 입력: (values: {customer, issue_date, from_month, to_month, items:[{name, spec, monthly_amount}]})
# 출력: (customer, issue: date, lo: (연, 월), hi: (연, 월), items: 앞뒤 공백을 지운 품목 목록)
# 처리: 거래처·작성 일자·청구 시작/끝 월·품목 1개 이상·각 품목 이름/규격(공백 제외 비지 않음)·월 금액(1~10^12 정수, bool 제외) 검사
#       → 시작 월 > 끝 월이면 거절
# 예외/실패: 하나라도 어기면 ValueError(한글 안내; 품목 오류는 'N번째 품목' 번호 포함)
# 부수효과: 없음
def validate(values):
    customer = str(values.get("customer") or "").strip()
    if not customer:
        raise ValueError("수신(거래처)을 입력해 주세요.")
    try:
        issue = date.fromisoformat(str(values.get("issue_date") or ""))
    except ValueError:
        raise ValueError("작성 일자를 YYYY-MM-DD 형식으로 입력해 주세요.")
    lo = parse_month(values.get("from_month"), "청구 시작 월")
    hi = parse_month(values.get("to_month"), "청구 끝 월")
    if lo > hi:
        raise ValueError("청구 시작 월이 끝 월보다 늦습니다.")
    items = []
    for i, it in enumerate(values.get("items") or [], start=1):
        name, spec, amount = str(it.get("name") or "").strip(), str(it.get("spec") or "").strip(), it.get("monthly_amount")
        if not name:
            raise ValueError(f"{i}번째 품목의 품목명을 입력해 주세요.")
        if not spec:
            raise ValueError(f"{i}번째 품목의 규격을 입력해 주세요.")
        if not is_amount(amount) or amount == 0:
            raise ValueError(f"{i}번째 품목의 월금액을 1원 이상의 정수로 입력해 주세요.")
        items.append({"name": name, "spec": spec, "monthly_amount": amount})
    if not items:
        raise ValueError("품목을 1개 이상 입력해 주세요.")
    return customer, issue, lo, hi, items


# [기능명세] output_name
# 목적: 결과 파일의 기본 이름(확장자 제외) 결정
# 입력: (contract_no: 관리번호 또는 None, customer: 거래처, lo/hi: 청구 범위 (연, 월), full: 계약기간 전체 범위인지)
# 출력: "거래명세서_{관리번호}_{거래처}[_{YYYYMM}-{YYYYMM}]" — files.safe_name으로 정리하고 연속 마침표는 하나로 줄인 문자열
# 처리: 관리번호가 없으면 생략; 계약기간 전체가 아니면(기간을 모르는 경우 포함) 범위 접미사를 붙임
# 예외/실패: 없음
# 부수효과: 없음
def output_name(contract_no, customer, lo, hi, full):
    stem = "_".join(x for x in ("거래명세서", contract_no, customer) if x)
    if not full:
        stem += f"_{lo[0]}{lo[1]:02d}-{hi[0]}{hi[1]:02d}"
    # 다운로드 API가 '..'이 든 이름을 거절하므로 연속 마침표를 미리 하나로 줄인다.
    return re.sub(r"\.{2,}", ".", safe_name(stem)) or "거래명세서"


# [기능명세] build
# 목적: 확인한 값으로 거래명세서 템플릿 사본을 채워 저장 폴더에 새 파일로 저장(스킬 fill_statement.py의 채우기·검사 규칙)
# 입력: (values: {customer, issue_date, from_month, to_month, items, contract_no?, period_start?, period_end?},
#        out_dir: 저장 폴더 Path, template=None: 템플릿 경로(기본 template_path()))
# 출력: {file: 파일명, path: 절대경로 문자열, rows: 채운 줄 수, supply: 공급가액 합계 int, vat: 세액(10%, 정수면 int), total: 합계, warnings: [str]}
# 처리: validate → 청구월×품목(월 바깥, 품목 안쪽) 줄 목록, 12줄 초과 거절 → 계약기간 월과 다르면 warnings
#       → 템플릿 열기 → 쓰기 전 전체 칸 값 기록 → A3 수신·A5 작성일 치환 → 13행부터 A=청구월 말일(date, yyyy-mm-dd), B=품목, C=규격,
#         D=월금액(int), E=청구월(범위가 한 해면 "1월", 해가 섞이면 "26.03") → 쓴 칸 외 모든 칸(F~H·25행 수식, 공급자 칸)이 그대로인지,
#         템플릿 자리표시자가 남았는지 검사 → fullCalcOnLoad=True → 같은 이름이 있으면 _2, _3…으로 새 이름을 골라 배타적 생성("xb")으로 저장
# 예외/실패: 필수값 누락·형식 오류·from>to·12줄 초과·자리표시자 없음/잔존·수식 변경은 ValueError(한글)이며 파일을 만들지 않음;
#           저장 중 오류면 만들던 파일을 지우고 예외 전파
# 부수효과: out_dir 생성과 새 xlsx 파일 1개 생성. 템플릿 원본과 기존 파일은 절대 덮어쓰지 않음
def build(values, out_dir, template=None):
    template = Path(template or template_path())
    customer, issue, lo, hi, items = validate(values)
    months = list(month_range(lo, hi))
    rows = [(y, m, it) for (y, m) in months for it in items]
    if len(rows) > CAPACITY:
        raise ValueError(f"청구월 {len(months)}개 × 품목 {len(items)}개 = {len(rows)}줄입니다. 템플릿은 {CAPACITY}줄까지만 넣을 수 있으니 청구 범위를 나눠 여러 번 만들어 주세요.")
    multi_year = len({y for y, _ in months}) > 1

    warnings = []
    start, end = month_of(values.get("period_start")), month_of(values.get("period_end"))
    full = False
    if start and end:
        period = (tuple(map(int, start.split("-"))), tuple(map(int, end.split("-"))))
        full = (lo, hi) == period
        if not full:
            warnings.append(f"청구 범위 {lo[0]}-{lo[1]:02d}~{hi[0]}-{hi[1]:02d}가 계약기간 {start}~{end}와 다릅니다.")

    wb = openpyxl.load_workbook(template)
    ws = wb.active
    before = {c.coordinate: c.value for row in ws.iter_rows() for c in row}
    a3, a5 = str(ws["A3"].value or ""), str(ws["A5"].value or "")
    if "{거래처}" not in a3 or not all(k in a5 for k in ("{YYYY}", "{MM}", "{DD}")):
        raise ValueError("템플릿 A3/A5에 {거래처}·{YYYY}·{MM}·{DD} 자리표시자가 없습니다. 템플릿 파일을 확인해 주세요.")
    a3_rest = a3.replace("{거래처}", "")
    a5_new = a5.replace("{YYYY}", f"{issue.year:04d}").replace("{MM}", f"{issue.month:02d}").replace("{DD}", f"{issue.day:02d}")
    ws["A3"] = safe_cell(a3.replace("{거래처}", customer))
    ws["A5"] = safe_cell(a5_new)
    written = {"A3", "A5"}
    for r, (y, m, it) in enumerate(rows, start=FIRST_ROW):
        ws[f"A{r}"] = date(y, m, calendar.monthrange(y, m)[1])
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
        ws[f"B{r}"] = safe_cell(it["name"])
        ws[f"C{r}"] = safe_cell(it["spec"])
        ws[f"D{r}"] = int(it["monthly_amount"])
        ws[f"E{r}"] = f"{y % 100:02d}.{m:02d}" if multi_year else f"{m}월"
        written.update(f"{col}{r}" for col in "ABCDE")

    errors = [f"{ref} 칸(수식·고정값)이 바뀌었습니다." for ref, value in before.items() if ref not in written and ws[ref].value != value]
    # 템플릿에서 온 글자에만 자리표시자 검사를 한다(사용자가 입력한 품목·규격·거래처의 중괄호는 허용).
    for ref, value in [("A3", a3_rest), ("A5", a5_new)] + [(ref, v) for ref, v in before.items() if ref not in written]:
        if isinstance(value, str) and not value.startswith("=") and re.search(r"\{[^{}]*\}", value):
            errors.append(f"{ref} 칸에 자리표시자가 남았습니다: {value}")
    if errors:
        raise ValueError("거래명세서를 저장하지 않았습니다. " + " ".join(errors))

    # openpyxl은 수식 계산값을 저장하지 않으므로 엑셀에서 열 때 다시 계산하게 한다.
    wb.calculation.fullCalcOnLoad = True
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = output_name(values.get("contract_no"), customer, lo, hi, full)
    n = 1
    while True:
        path = out_dir / (stem + (f"_{n}" if n > 1 else "") + ".xlsx")
        if path.resolve() == template.resolve():
            n += 1
            continue
        try:
            # "xb": 같은 이름 파일이 있으면 실패하므로 기존 파일을 절대 덮어쓰지 않는다.
            handle = open(path, "xb")
        except FileExistsError:
            n += 1
            continue
        try:
            with handle:
                wb.save(handle)
        except BaseException:
            path.unlink(missing_ok=True)
            raise
        break

    supply = sum(it["monthly_amount"] for _, _, it in rows)
    # 템플릿 G열은 F*0.1(반올림 없음)이므로 같은 값을 Decimal로 정확히 계산한다.
    vat_dec = Decimal(supply) * Decimal("0.1")
    vat = int(vat_dec) if vat_dec == vat_dec.to_integral_value() else float(vat_dec)
    total = int(supply + vat_dec) if isinstance(vat, int) else float(supply + vat_dec)
    return {"file": path.name, "path": str(path.resolve()), "rows": len(rows), "supply": supply, "vat": vat, "total": total, "warnings": warnings}
