# [파일 기능명세] 전송 전 문서 편집 3단 화면(거래명세서·지출결의서 확인 창)을 실제 Chromium으로 검증한다.
# 입력: 빌드된 frontend/dist, 설치된 Playwright Chromium. verify_statement.py의 시드·서버 시작·계약 선택 함수를 재사용한다.
# 출력: tmp/docedit-review/<시각>/checks.json·PNG·server.log, 요약 JSON(표준 출력), 종료코드(실패 있으면 1).
# 처리: 격리 데이터(계약 2건)로 서버 실행 → 로그인 → 계약서 목록
#   → 거래명세서(C-2026-002): 3단 구성·목차 3개, 문서 편집→마크다운 반영, 마크다운 편집→문서·미리보기 반영, 형식 오류 문구와 버튼 비활성,
#     닫기 확인 줄(DiscardBar), 편집 값으로 생성한 xlsx 내용 확인
#   → 지출결의서(C-2026-011): 목차 #~### 만 표시·목차 클릭 이동, 상신 요청 본문(POST를 가로채 서버로 보내지 않음) 확인
#   → 390px: 탭 3개 전환·가로 넘침 없음 → JavaScript 오류 0건.
# 부수효과: tmp/docedit-review 아래 파일 생성, 테스트 서버 실행·종료. 실제 사이트·사용자 DB에는 접속·변경하지 않는다.
#   지출결의서 상신 POST는 브라우저에서 가짜 응답으로 대체하므로 그룹웨어에 기안이 만들어지지 않는다.
import json
from datetime import datetime
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT/'scripts'))
from playwright.sync_api import sync_playwright, expect

import verify_statement as vs


# [기능명세] set_label
# 목적: 마크다운 원문에서 '- 라벨: 값' 줄의 값을 바꾼다(없으면 변경하지 않는다).
# 입력: (md: 마크다운 문자열, name: 라벨, value: 새 값)
# 출력: 바뀐 마크다운 문자열
# 처리: 줄 시작의 '- name:' 한 줄을 정규식으로 교체
# 예외/실패: 없음. 부수효과: 없음
def set_label(md, name, value):
    return re.sub(r'^- ' + re.escape(name) + r':.*$', lambda _m: f'- {name}: {value}', md, count=1, flags=re.M)


# [기능명세] main
# 목적: 위 처리 순서대로 문서 편집 화면을 검증하고 결과를 저장한다.
# 입력: 없음. 출력: 종료코드. 처리/예외: 단계별 예외는 실패 검사로 기록하고 계속 진행. 시드·서버 실패는 즉시 종료.
# 부수효과: tmp/docedit-review 산출물, 테스트 서버·브라우저(finally에서 종료)
def main():
    out = ROOT/'tmp/docedit-review'/datetime.now().strftime('%Y%m%d-%H%M%S')
    out.mkdir(parents=True)
    data, pdf_root = out/'data', out/'pdf-root'
    checks, errors, server = vs.Checks(), [], None
    log = (out/'server.log').open('w', encoding='utf-8')
    try:
        if checks.step('격리 데이터 준비', lambda: vs.seed_data(data, pdf_root) or True) is None:
            return vs.finish(out, checks)
        started = checks.step('테스트 서버 시작', lambda: vs.start_server(data, log))
        if not started:
            return vs.finish(out, checks)
        server, base, token = started
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                page = browser.new_page(viewport={'width': 1440, 'height': 960}, accept_downloads=True)
                page.on('pageerror', lambda error: errors.append(str(error)))

                def login():
                    page.goto(base+'/?token='+token)
                    page.get_by_label('비밀번호', exact=True).fill('0000')
                    page.get_by_role('button', name='시작하기', exact=True).click()
                    page.get_by_role('button', name='계약서 목록', exact=True).click()
                    expect(page.get_by_role('button', name=re.compile('C-2026-002')).first).to_be_visible(timeout=10000)
                    return True
                if not checks.step('로그인·계약서 목록', login):
                    return vs.finish(out, checks)

                # ── 거래명세서 ──
                def statement_edit():
                    vs.open_contract(page, 'C-2026-002')
                    vs.statement_button(page).click()
                    dialog = page.get_by_role('dialog')
                    expect(dialog).to_be_visible(timeout=10000)
                    md_pane = dialog.get_by_role('region', name='마크다운 편집')
                    doc_pane = dialog.get_by_role('region', name='문서 편집')
                    toc = dialog.get_by_role('navigation', name='목차')
                    textarea = dialog.get_by_label('마크다운 원문')
                    for name, pane in (('마크다운 편집', md_pane), ('문서 편집', doc_pane), ('목차', toc)):
                        checks.add(f'3단 {name} 표시', pane.count() == 1 and pane.is_visible())
                    items = toc.get_by_role('button')
                    checks.add('목차 3개(거래명세서/기본 정보/품목)', items.all_inner_texts() == ['거래명세서', '기본 정보', '품목'], items.all_inner_texts())
                    checks.add('문서 편집에 표 표시', doc_pane.locator('table').count() == 1 and doc_pane.locator('th').count() == 3)

                    # 문서(중) 편집 → 마크다운(좌) 반영
                    doc_pane.get_by_text('수신: 세움바이오(주)').click()
                    page.keyboard.press('End')
                    page.keyboard.type('(WY)')
                    expect(textarea).to_have_value(re.compile(r'- 수신: 세움바이오\(주\)\(WY\)'), timeout=5000)
                    checks.add('문서 편집이 마크다운에 반영', True)

                    # 마크다운(좌) 편집 → 문서(중)·미리보기 반영: 월 금액 190000 → 200000
                    textarea.fill(textarea.input_value().replace('190000', '200000'))
                    expect(doc_pane.get_by_text('200000')).to_be_visible(timeout=5000)
                    checks.add('마크다운 편집이 문서에 반영', True)
                    expect(dialog).to_contain_text('2,400,000원', timeout=5000)
                    text = dialog.inner_text()
                    checks.add('미리보기가 편집 값으로 갱신', all(v in text for v in ('2,400,000', '240,000', '2,640,000')))
                    page.screenshot(path=str(out/'statement-edit-desktop.png'), full_page=True)

                    # 형식 오류: 수신 줄 삭제 → 문구·버튼 비활성, 복원하면 해소
                    original = textarea.input_value()
                    textarea.fill(re.sub(r'^- 수신:.*\n', '', original, flags=re.M))
                    problems = dialog.get_by_label('생성 전에 확인할 것')
                    expect(problems).to_contain_text("'수신' 항목이 없습니다", timeout=5000)
                    checks.add('수신 누락 문구', True)
                    checks.add('수신 누락 시 생성 버튼 비활성', dialog.get_by_role('button', name='거래명세서 생성', exact=True).is_disabled())
                    textarea.fill(original)
                    expect(dialog.get_by_role('button', name='거래명세서 생성', exact=True)).to_be_enabled(timeout=5000)
                    checks.add('복원하면 생성 버튼 활성', True)

                    # 닫기 확인 줄: Esc → 확인 줄, 계속 편집, 다시 Esc → 버리고 닫기
                    page.keyboard.press('Escape')
                    expect(dialog.get_by_test_id('doc-discard')).to_be_visible(timeout=5000)
                    checks.add('변경 후 Esc는 닫지 않고 확인 줄 표시', dialog.count() == 1)
                    dialog.get_by_role('button', name='계속 편집').click()
                    expect(dialog.get_by_test_id('doc-discard')).to_have_count(0)
                    checks.add('계속 편집 선택 시 확인 줄 사라짐', dialog.count() == 1 and textarea.input_value() == original)
                    return dialog
                dialog = checks.step('거래명세서 편집 화면', statement_edit)

                def statement_generate():
                    dialog.get_by_role('button', name='거래명세서 생성', exact=True).click()
                    expect(dialog.get_by_role('link', name='내려받기')).to_be_visible(timeout=20000)
                    files = sorted((pdf_root/'거래명세서').glob('*.xlsx'))
                    checks.add('편집 후 생성한 xlsx 1개', len(files) == 1, [f.name for f in files])
                    import openpyxl
                    ws = openpyxl.load_workbook(files[0]).active
                    checks.add('xlsx 수신에 문서 편집 반영', '세움바이오(주)(WY)' in str(ws['A3'].value), ws['A3'].value)
                    amounts = [ws[f'D{r}'].value for r in range(13, 25)]
                    checks.add('xlsx 월 금액에 마크다운 편집 반영', amounts == [200000] * 12, amounts)
                    checks.add('xlsx 품목명 유지', ws['B13'].value == vs.EXPECT['item_name'], ws['B13'].value)
                    dialog.get_by_role('button', name='닫기', exact=True).click()
                    expect(page.get_by_role('dialog')).to_have_count(0)
                    return True
                if dialog is not None:
                    checks.step('편집 값으로 거래명세서 생성', statement_generate)

                # 버리고 닫기: 새 창에서 변경 후 Esc → 버리고 닫기
                def discard_close():
                    vs.open_contract(page, 'C-2026-002')
                    vs.statement_button(page).click()
                    dlg = page.get_by_role('dialog')
                    expect(dlg).to_be_visible(timeout=10000)
                    area = dlg.get_by_label('마크다운 원문')
                    area.fill(area.input_value() + '\n')
                    page.keyboard.press('Escape')
                    expect(dlg.get_by_test_id('doc-discard')).to_be_visible(timeout=5000)
                    dlg.get_by_role('button', name='버리고 닫기').click()
                    expect(page.get_by_role('dialog')).to_have_count(0)
                    checks.add('버리고 닫기로 창이 닫힘', True)
                    return True
                checks.step('버리고 닫기', discard_close)

                # ── 지출결의서 ──
                captured = []

                def expense_edit():
                    # [콜백 명세] 상신 POST만 가로채 본문을 저장하고 가짜 접수 응답을 준다(GET은 그대로 통과).
                    def handler(route):
                        if route.request.method == 'POST':
                            captured.append(route.request.post_data_json)
                            route.fulfill(status=200, content_type='application/json', body='{"id":"job-docedit-test"}')
                        else:
                            route.continue_()
                    page.route(re.compile(r'/api/contracts/C-2026-011/expense$'), handler)
                    detail = vs.open_contract(page, 'C-2026-011')
                    detail.get_by_test_id('expense-open').click()
                    modal = page.get_by_test_id('expense-modal')
                    expect(modal).to_be_visible(timeout=10000)
                    textarea = modal.get_by_label('마크다운 원문')
                    toc = modal.get_by_role('navigation', name='목차')
                    md = textarea.input_value()
                    labels = [name for name in ('제목', '청구·사용 월', '공급가액', '세액', '합계') if re.search(r'^- ' + name + ':', md, re.M)]
                    checks.add('지출결의서 문서 라벨 5개', labels == ['제목', '청구·사용 월', '공급가액', '세액', '합계'], labels)
                    checks.add('지출결의서 섹션 구성', all(h in md for h in ('# 지출결의서', '## 기본 정보', '## 금액', '## 첨부', '## 특이사항')))

                    # 금액·월을 채우고 긴 특이사항(제목 #~#### 포함)을 넣는다.
                    md = set_label(md, '청구·사용 월', '2026-09')
                    md = set_label(md, '공급가액', '2100000')
                    md = set_label(md, '세액', '210000')
                    md = set_label(md, '합계', '2310000')
                    head = md[:md.index('## 특이사항')]
                    note = '메모 시작\n\n' + '\n\n'.join(f'본문 줄 {i}' for i in range(80)) + '\n\n### 마지막 제목\n\n#### 제외될 제목\n\n끝'
                    textarea.fill(head + '## 특이사항\n\n' + note + '\n')
                    expect(modal.get_by_role('button', name='상신')).to_be_enabled(timeout=5000)
                    checks.add('채운 뒤 상신 버튼 활성', True)
                    titles = toc.get_by_role('button').all_inner_texts()
                    checks.add('목차에 # ~ ### 만(#### 제외)', '마지막 제목' in titles and '제외될 제목' not in titles, titles)

                    # 목차 클릭 → 좌 편집기가 아래로 스크롤, 중 편집기에서 해당 제목이 보임
                    toc.get_by_role('button', name='기본 정보').click()
                    page.wait_for_timeout(300)
                    before = textarea.evaluate('el => el.scrollTop')
                    checks.add('목차 클릭으로 위쪽 제목 이동(스크롤 위)', before < 100, before)
                    toc.get_by_role('button', name='마지막 제목').click()
                    page.wait_for_timeout(500)
                    after = textarea.evaluate('el => el.scrollTop')
                    checks.add('목차 클릭으로 마크다운 편집기 스크롤', after > before, {'before': before, 'after': after})
                    heading = modal.get_by_role('region', name='문서 편집').get_by_role('heading', name='마지막 제목')
                    inside = heading.evaluate("""el => {const pane = el.closest('.doc-pane'); const a = el.getBoundingClientRect(), b = pane.getBoundingClientRect();
                        return a.top >= b.top - 1 && a.bottom <= b.bottom + 1;}""")
                    checks.add('목차 클릭으로 문서 편집기 제목이 보임', inside)
                    page.screenshot(path=str(out/'expense-edit-desktop.png'), full_page=True)

                    # 상신 → 가로챈 요청 본문 확인(가짜 접수 응답이라 실제 상신은 없다)
                    modal.get_by_test_id('expense-submit').click()
                    expect(modal.get_by_role('status')).to_be_visible(timeout=10000)
                    body = captured[0] if captured else {}
                    checks.add('상신 요청 1회', len(captured) == 1, len(captured))
                    checks.add('상신 본문 금액·월', (body.get('supply'), body.get('vat'), body.get('total'), body.get('month')) == (2100000, 210000, 2310000, '2026-09'), body)
                    checks.add('상신 본문 제목·첨부', bool(body.get('title')) and isinstance(body.get('attachments'), list) and len(body['attachments']) >= 1, body)
                    checks.add('상신 본문 특이사항에 편집 내용', '메모 시작' in body.get('note', '') and '### 마지막 제목' in body.get('note', '') and '끝' in body.get('note', ''), body.get('note', '')[:80])
                    return True
                checks.step('지출결의서 편집 화면', expense_edit)
                page.unroute(re.compile(r'/api/contracts/C-2026-011/expense$'))

                # ── 390px: 탭 전환 ──
                def mobile():
                    page.reload()
                    page.get_by_role('button', name='계약서 목록', exact=True).click()
                    page.set_viewport_size({'width': 390, 'height': 844})
                    page.wait_for_timeout(300)
                    vs.open_contract(page, 'C-2026-002')
                    button = vs.statement_button(page)
                    button.scroll_into_view_if_needed()
                    button.click()
                    dlg = page.get_by_role('dialog')
                    expect(dlg).to_be_visible(timeout=10000)
                    tabs = dlg.get_by_role('tab')
                    checks.add('390px 탭 3개', tabs.all_inner_texts() == ['마크다운', '문서', '목차'], tabs.all_inner_texts())
                    for name, shown in (('마크다운', '마크다운 원문'), ('문서', '문서 본문'), ('목차', None)):
                        dlg.get_by_role('tab', name=name).click()
                        page.wait_for_timeout(150)
                        if shown == '마크다운 원문':
                            checks.add('탭 마크다운: 원문 표시', dlg.get_by_label('마크다운 원문').is_visible())
                        elif shown == '문서 본문':
                            checks.add('탭 문서: 본문 표시·원문 숨김', dlg.get_by_role('textbox', name='문서 본문').is_visible() and not dlg.get_by_label('마크다운 원문').is_visible())
                        else:
                            checks.add('탭 목차: 목차 항목 표시', dlg.get_by_role('tabpanel', name='목차').get_by_role('button').count() == 3 if dlg.get_by_role('tabpanel', name='목차').count() else dlg.get_by_role('button', name='품목').is_visible())
                    overflow = page.evaluate("()=>document.documentElement.scrollWidth-document.documentElement.clientWidth")
                    checks.add('390px 페이지 가로 넘침 없음', overflow <= 1, overflow)
                    box = dlg.bounding_box()
                    checks.add('390px 편집 창이 화면 폭 안', box is not None and box['x'] >= -1 and box['x'] + box['width'] <= 391, box)
                    page.screenshot(path=str(out/'statement-edit-390.png'))
                    return True
                checks.step('390px 탭 화면', mobile)

                checks.add('JavaScript 오류 0건', not errors, errors)
            finally:
                browser.close()
    finally:
        if server:
            server.terminate()
            server.wait(timeout=20)
        log.close()
    return vs.finish(out, checks)


if __name__ == '__main__':
    sys.stdout.reconfigure(errors='replace')
    sys.exit(main())
