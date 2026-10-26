# [파일 기능명세]
# 리팩토링 2단계(도메인) 회귀 검사. 다음을 검증한다.
#  (a) store.py가 service를 import하지 않는다(순환 제거) — 새 프로세스에서 import saebom.store만 하고 sys.modules를 확인.
#  (b) document_mover 함수 단위: 경로 검증 실패, 복사본 해시 불일치 시 생성 파일 기록, 같은 이름 다른 내용 충돌, 이전 사본 정리.
#  (c) results.summary 열: 구스키마 DB를 열면 열이 추가되고 첫 result_summaries가 요약을 채우며, 새 결과는 summary가 함께 저장된다.
#  (d) Service.get_result/list_results/document_file(없는 문서·없는 파일은 None).
#  (e) spreadsheet_safe 통일 후 거래명세서 셀이 탭·캐리지 리턴 접두까지 막는다.
#  (f) 명시 import 후에도 saebom.service.extract monkeypatch가 로컬 세금계산서 판독 캐시에 적용된다.
# 외부 네트워크·브라우저·사용자 앱 데이터를 쓰지 않으며 모든 DB·파일은 tmp_path 안에 만든다.
# 유지보수: store.py·document_mover.py·services/*의 계약이 바뀌면 이 검사와 기능명세를 함께 갱신한다.

import hashlib
import json
import sqlite3
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from saebom import document_mover
from saebom.service import Service
from saebom.statement import build, safe_cell
from saebom.store import Store


# [기능명세] test_store_import_does_not_load_service
# 목적: R2-1 순환 제거 검증
# 처리: 새 파이썬 프로세스에서 saebom.store를 import한 뒤 saebom.service가 sys.modules에 없는지, 정렬 키가 동작하는지 확인
# 예외/실패: 단언 실패
# 부수효과: 하위 프로세스 1개
def test_store_import_does_not_load_service(tmp_path):
    code = ("import sys; import saebom.store as s; "
            "assert 'saebom.service' not in sys.modules, 'service가 로드됨'; "
            f"st = s.Store(r'{tmp_path}'); "
            "st.save_contract_import({'id': 'C-2026-002'}); st.save_contract_import({'id': 'C-2026-010'}); "
            "assert [c['id'] for c in st.contract_imports()] == ['C-2026-010', 'C-2026-002']; "
            "assert 'saebom.service' not in sys.modules; st.close()")
    root = Path(__file__).resolve().parent.parent
    done = subprocess.run([sys.executable, "-X", "utf8", "-c", code], cwd=root, capture_output=True, text=True, encoding="utf-8")
    assert done.returncode == 0, done.stderr


# [기능명세] make_doc
# 목적: document_mover 시험용 원본 파일과 문서 사전 생성
# 입력: (folder, name, body: bytes, kind)
# 출력: (문서 사전, 원본 Path)
# 부수효과: folder에 파일 생성
def make_doc(folder, name, body=b"%PDF-1.4 test", kind="invoice"):
    folder.mkdir(parents=True, exist_ok=True)
    src = folder / name
    src.write_bytes(body)
    return {"id": "d1", "source": "http://x/" + name, "name": name, "hash": hashlib.sha256(body).hexdigest(), "kind": kind}, src.resolve()


# [기능명세] test_prepare_root_rejects_bad_paths
# 목적: 상대경로·UNC 경로 거절과 하위 폴더 준비, 시험 파일 미잔류 검증
def test_prepare_root_rejects_bad_paths(tmp_path):
    with pytest.raises(ValueError, match="절대 폴더"):
        document_mover.prepare_root("상대/경로")
    with pytest.raises(ValueError, match="네트워크 공유"):
        document_mover.prepare_root("\\\\server\\share")
    root = document_mover.prepare_root(tmp_path / "new")
    assert (root / "세금계산서").is_dir() and (root / "계약서").is_dir()
    assert not [p for p in root.iterdir() if p.name.startswith(".saebom-write-")]


# [기능명세] test_copy_document_success_and_same_path
# 목적: 정상 복사(해시 검증·created 기록)와 이미 대상 위치인 문서의 무복사 반환 검증
def test_copy_document_success_and_same_path(tmp_path):
    doc, src = make_doc(tmp_path / "old", "세금계산서_a.pdf")
    root = document_mover.prepare_root(tmp_path / "new")
    created = []
    doc_id, source, target = document_mover.copy_document(doc, src, root, set(), set(), created)
    assert doc_id == "d1" and source == src and created == [target]
    assert target.parent.name == "세금계산서" and target.read_bytes() == src.read_bytes() and src.exists()
    again = []
    assert document_mover.copy_document(doc, target, root, set(), set(), again)[2] == target and again == []


# [기능명세] test_copy_document_failures_keep_state
# 목적: 원본 누락·원본 변조·같은 이름 다른 내용 충돌이 ValueError이며 기존 파일을 보존하는지 검증
def test_copy_document_failures_keep_state(tmp_path):
    doc, src = make_doc(tmp_path / "old", "계약_a.pdf", kind="contract")
    root = document_mover.prepare_root(tmp_path / "new")
    with pytest.raises(ValueError, match="원본 파일을 찾을 수 없어"):
        document_mover.copy_document(doc, tmp_path / "none.pdf", root, set(), set(), [])
    with pytest.raises(ValueError, match="원본 파일 내용이 달라"):
        document_mover.copy_document({**doc, "hash": "0" * 64}, src, root, set(), set(), [])
    clash = root / "계약서" / (doc["hash"][:12] + "_" + doc["name"])
    clash.write_bytes("다른 내용".encode())
    created = []
    with pytest.raises(ValueError, match="내용이 다른 같은 이름"):
        document_mover.copy_document(doc, src, root, set(), set(), created)
    assert clash.read_bytes() == "다른 내용".encode() and created == []


# [기능명세] test_copy_document_hash_mismatch_recorded_for_cleanup
# 목적: 복사본 해시가 어긋나면 OSError를 내되 생성한 파일이 created에 남아 호출자가 지울 수 있는지 검증
def test_copy_document_hash_mismatch_recorded_for_cleanup(tmp_path, monkeypatch):
    doc, src = make_doc(tmp_path / "old", "세금계산서_b.pdf")
    root = document_mover.prepare_root(tmp_path / "new")
    monkeypatch.setattr(document_mover.shutil, "copyfileobj", lambda a, b: b.write(b"broken"))
    created = []
    with pytest.raises(OSError, match="일치하지 않습니다"):
        document_mover.copy_document(doc, src, root, set(), set(), created)
    assert len(created) == 1 and created[0].exists() and src.exists()


# [기능명세] test_remove_old_copies_protects_destinations
# 목적: 이전 사본만 지우고 다른 이동의 대상 경로는 보호하는지 검증
def test_remove_old_copies_protects_destinations(tmp_path):
    a, b, c = (tmp_path / n for n in "abc")
    for p in (a, b, c):
        p.write_bytes(b"x")
    document_mover.remove_old_copies([("1", a, b), ("2", b, c), ("3", c, c)])
    assert not a.exists() and b.exists() and c.exists()


# [기능명세] test_kind_folder_guess
# 목적: kind가 없을 때 기존 추측 규칙이 유지되는지 검증
def test_kind_folder_guess(tmp_path):
    src = tmp_path / "아무폴더" / "f.pdf"
    assert document_mover.kind_folder({"kind": "contract", "source": "s", "name": "n"}, src, set(), set()) == "계약서"
    assert document_mover.kind_folder({"source": "s", "name": "세금계산서_x"}, src, set(), set()) == "세금계산서"
    assert document_mover.kind_folder({"source": "u", "name": "n"}, src, set(), {"u"}) == "계약서"
    assert document_mover.kind_folder({"source": "s", "name": "n"}, src, set(), set()) == "계약서"


# [기능명세] test_store_move_uses_mover
# 목적: Store.set_document_root가 mover를 거쳐 이동·DB 확정·이전 사본 삭제를 끝까지 수행하는지 검증
def test_store_move_uses_mover(tmp_path):
    store = Store(tmp_path / "data")
    doc, src = make_doc(tmp_path / "old", "세금계산서_c.pdf")
    store.save_document(doc["source"], doc["name"], doc["hash"], {"fields": {}}, src, kind="invoice")
    new = store.set_document_root(tmp_path / "new")
    moved = store.document(source=doc["source"])
    assert Path(moved["disk_path"]).parent == Path(new) / "세금계산서" and Path(moved["disk_path"]).is_file()
    assert not src.exists() and store.document_root() == Path(new)
    store.close()


OLD_SCHEMA = """
CREATE TABLE items(kind TEXT, id TEXT, data TEXT, PRIMARY KEY(kind,id));
CREATE TABLE documents(id TEXT PRIMARY KEY, source TEXT UNIQUE, name TEXT, hash TEXT, extraction TEXT, disk_path TEXT, kind TEXT);
CREATE TABLE results(id TEXT PRIMARY KEY, created TEXT, data TEXT);
CREATE TABLE meta(key TEXT PRIMARY KEY, value TEXT);
"""


# [기능명세] test_old_schema_results_get_summary
# 목적: summary 열이 없는 구스키마 DB를 열면 열이 추가되고, 첫 result_summaries가 요약을 채우며 둘째 조회가 같은 결과를 주는지 검증
# 처리: tmp_path에 구스키마 DB 생성(결과 3건, created 동률 포함) → Store 열기 → 열 확인 → 요약 비교(순서 포함) → data를 망가뜨려도 재조회 성공
def test_old_schema_results_get_summary(tmp_path):
    root = tmp_path / "data"
    root.mkdir()
    db = sqlite3.connect(root / "saebom.sqlite3")
    db.executescript(OLD_SCHEMA)
    results = [
        {"id": "r1", "created": "2026-10-01T09:00:00+09:00", "rows": [{"mail_id": "m1", "partner": "가", "status": "불일치"}, {"mail_id": "m2", "partner": "나", "status": "일치"}]},
        {"id": "r2", "created": "2026-10-02T09:00:00+09:00", "parent": "r1", "rows": []},
        {"id": "r3", "created": "2026-10-02T09:00:00+09:00", "rows": [{"mail_id": "m9", "partner": "다", "status": "불일치"}] * 5},
    ]
    for r in results:
        db.execute("INSERT INTO results VALUES(?,?,?)", (r["id"], r["created"], json.dumps(r, ensure_ascii=False)))
    db.commit()
    db.close()

    store = Store(root)
    assert "summary" in {r[1] for r in store.db.execute("PRAGMA table_info(results)")}
    assert store.db.execute("SELECT COUNT(*) FROM results WHERE summary IS NULL").fetchone()[0] == 3
    first = store.result_summaries()
    assert [s["id"] for s in first] == ["r3", "r2", "r1"]
    assert first[0] == {"id": "r3", "created": results[2]["created"], "parent": None, "count": 5, "mismatch": 5,
                        "partners": ["다", "다", "다"], "mail_ids": ["m9"] * 5, "verdicts": first[0]["verdicts"]}
    assert [v["mail_id"] for v in first[0]["verdicts"]] == ["m9"] * 5
    assert first[1]["parent"] == "r1" and first[1]["count"] == 0 and first[1]["partners"] == []
    assert first[2]["mismatch"] == 1 and first[2]["mail_ids"] == ["m1", "m2"]
    assert store.db.execute("SELECT COUNT(*) FROM results WHERE summary IS NULL").fetchone()[0] == 0
    # 이후 조회는 data를 읽지 않고도 같은 결과여야 한다: data를 망가뜨려 확인한다.
    store.db.execute("UPDATE results SET data='broken'")
    store.db.commit()
    assert store.result_summaries() == first
    store.close()
    # 다시 열어도(열이 이미 있음) 오류 없이 동일하다.
    store = Store(root)
    assert store.result_summaries() == first
    store.close()


# [기능명세] test_save_result_stores_summary
# 목적: 새 결과 저장 시 summary가 함께 저장되고 result_summaries가 그 값을 쓰는지 검증
def test_save_result_stores_summary(tmp_path):
    store = Store(tmp_path / "data")
    saved = store.save_result({"rows": [{"mail_id": "m1", "partner": "가", "status": "불일치"}, {"mail_id": "m2", "partner": "나", "status": "일치"}], "parent": None})
    raw = store.db.execute("SELECT summary FROM results WHERE id=?", (saved["id"],)).fetchone()[0]
    assert json.loads(raw) == {"id": saved["id"], "created": saved["created"], "parent": None, "count": 2, "mismatch": 1,
                               "partners": ["가", "나"], "mail_ids": ["m1", "m2"],
                               "verdicts": [{"mail_id": "m1", "contract_id": None, "status": "불일치"}, {"mail_id": "m2", "contract_id": None, "status": "일치"}]}
    assert store.result_summaries() == [json.loads(raw)]
    assert store.result(saved["id"])["rows"][0]["partner"] == "가"
    store.close()


# [기능명세] test_service_result_and_document_helpers
# 목적: Service.get_result/list_results/document_file 계약 검증(없는 ID·없는 파일은 None)
def test_service_result_and_document_helpers(tmp_path):
    store = Store(tmp_path / "data")
    service = Service(store)
    try:
        assert service.get_result("없음") is None and service.list_results() == []
        saved = store.save_result({"rows": [], "parent": None})
        assert service.get_result(saved["id"])["id"] == saved["id"]
        assert [s["id"] for s in service.list_results()] == [saved["id"]]

        pdf = tmp_path / "a.pdf"
        pdf.write_bytes(b"%PDF")
        doc = store.save_document("http://x/a", "a.pdf", "h" * 64, {"fields": {}}, pdf, kind="invoice")
        path, name = service.document_file(doc["id"])
        assert path == pdf and name == "a.pdf"
        assert service.document_file("없는문서") is None
        pdf.unlink()
        assert service.document_file(doc["id"]) is None
    finally:
        service.close()
        store.close()


# [기능명세] test_statement_cells_block_tab_and_cr_prefix
# 목적: spreadsheet_safe 통일 후 safe_cell이 탭·캐리지 리턴 접두까지 막고 build의 수식 이스케이프가 유지되는지 검증
def test_statement_cells_block_tab_and_cr_prefix(tmp_path):
    assert safe_cell("\t=1") == "'\t=1" and safe_cell("\r=1") == "'\r=1" and safe_cell("=1") == "'=1" and safe_cell("정상") == "정상"
    # build는 입력을 strip하므로 앞쪽 탭·CR은 셀에 닿기 전에 사라진다. 대신 안쪽에서 접두가 되는 값은 safe_cell이 막고, build의 수식 접두는 그대로 막힌다.
    values = {"customer": "거래처", "issue_date": "2026-10-01", "from_month": "2026-10", "to_month": "2026-10",
              "items": [{"name": "=cmd", "spec": "@x", "monthly_amount": 10}]}
    ws = openpyxl.load_workbook(build(values, tmp_path)["path"]).active
    assert ws["B13"].value == "'=cmd" and ws["C13"].value == "'@x" and ws["B13"].data_type == "s"


# [기능명세] test_shared_rules_keep_messages
# 목적: statement·expense가 rules로 판단을 옮긴 뒤에도 오류 문구와 상수 별칭이 유지되는지 검증
def test_shared_rules_keep_messages():
    from saebom import expense, rules, statement
    assert statement.MAX_AMOUNT == rules.MAX_AMOUNT == 10 ** 12
    for bad in ("2026-13", "2026-1", None, ""):
        with pytest.raises(ValueError, match="청구 시작 월을 YYYY-MM 형식으로"):
            statement.parse_month(bad, "청구 시작 월")
    assert statement.parse_month("2026-12", "x") == (2026, 12)
    base = {"title": "제목", "month": "2026-10", "supply": 100, "vat": 10, "total": 110}
    assert expense.validate(base)["month"] == "2026-10"
    with pytest.raises(ValueError, match="올바른 YYYY-MM"):
        expense.validate({**base, "month": "2026-13"})
    with pytest.raises(ValueError, match="YYYY-MM 형식으로 입력"):
        expense.validate({**base, "month": "26-10"})
    for bad in (True, 1.5, -1, 10 ** 12 + 1):
        with pytest.raises(ValueError, match="0~1조 원 범위"):
            expense.validate({**base, "supply": bad})


# [기능명세] test_explicit_imports_keep_service_namespace
# 목적: service.py 명시 import 후에도 기존 공개 이름이 남아 있고, service.extract monkeypatch가 로컬 판독 캐시에 먹히는지 검증
def test_explicit_imports_keep_service_namespace(tmp_path, monkeypatch):
    import saebom.service as service_module
    for name in ("Service", "extract", "allowed_url", "SiteClient", "Store", "contract_sort_key", "select_contracts", "LOCAL_FOLDER_KEY",
                 "LOCAL_AUDIT_KEY", "EMPTY_CONTRACTS_MESSAGE", "INVOICE_CACHE_LIMIT", "BusinessError", "ContractNotFound", "SubmitUncertain",
                 "strip_doc_pages", "JobRunner", "check_lists", "request_pdf", "MAX_PAIRS", "KIND_FOLDERS", "STATEMENT_FOLDER",
                 "safe_name", "compare", "norm", "now", "CONTRACT_ID"):
        assert hasattr(service_module, name), name

    store = Store(tmp_path / "data")
    service = Service(store)
    calls = []
    monkeypatch.setattr(service_module, "extract", lambda path, kind: calls.append((path, kind)) or {"fields": {"biz": {"value": "1"}}, "error": None, "pages": ["큼"]})
    pdf = tmp_path / "세금계산서.pdf"
    pdf.write_bytes(b"%PDF")
    try:
        first = service._cached_invoice(pdf, "invoice")
        second = service._cached_invoice(pdf, "invoice")
        assert len(calls) == 1 and first == second == {"fields": {"biz": {"value": "1"}}, "error": None}
        first["fields"]["biz"]["value"] = "변경"
        assert service._cached_invoice(pdf, "invoice")["fields"]["biz"]["value"] == "1"
    finally:
        service.close()
        store.close()
