# [파일 기능명세]
# 계약서 → 거래명세서 생성 기능 회귀 검사. statement.draft(fixture 계약서 001·002·003·011 초안 값, 여러 품목·총액·부가세 포함 안내),
# statement.build(12줄 채우기, A~E 값, F~H·25행 수식 보존, A3·A5 치환, 해가 바뀌는 청구월 표기, 12줄 초과 거절, 같은 이름 _2,
# 템플릿 원본 해시 불변, 수식 주입 방지)와 HTTP API(초안·생성·내려받기·지출 400·없는 계약 404·경로 조작 404·작업 중 400)를 검증한다.
# 외부 네트워크와 브라우저를 쓰지 않는다. 데이터 폴더는 Store(tmp_path/"data")로 격리하고 결과 파일은 tmp_path 안에만 만든다.
# 유지보수: statement.py 규칙이나 공유 API 형식(기획/거래명세서_생성_구축계획.md)이 바뀌면 이 검사와 기능명세를 함께 갱신한다.

import hashlib
from datetime import date
from pathlib import Path

from fastapi.testclient import TestClient
import openpyxl
import pytest

from saebom import config
from saebom.api import create_app
from saebom.compare import extract
from saebom.statement import build, draft, template_path
from saebom.store import Store

FIXTURES = Path(__file__).parent / "fixtures"
BASE = "http://127.0.0.1:12345"
FILE_BASE = "https://sbgw.vercel.app/files/"


# [기능명세] fixture_pdf
# 목적: 관리번호 끝 세 자리로 fixture 계약서 PDF 경로 찾기
# 입력: (no: "002" 같은 번호 문자열)
# 출력: Path
# 처리: tests/fixtures에서 계약서_C-2026-{no}_*.pdf 첫 파일
# 예외/실패: 없으면 StopIteration
# 부수효과: 파일 목록 읽기
def fixture_pdf(no):
    return next(FIXTURES.glob(f"계약서_C-2026-{no}_*.pdf"))


# [기능명세] contract_row
# 목적: contract_imports 계약 행 형태의 테스트 사전 생성(test_contracts_import.row와 같은 형식)
# 입력: (no: 번호 문자열, kind: "수입"|"지출")
# 출력: {id, kind, partner, title, period, amount_text, file:{name, url}}
# 처리: fixture 파일 이름으로 FILE_BASE URL을 만듦
# 예외/실패: 없음
# 부수효과: 없음
def contract_row(no, kind):
    name = fixture_pdf(no).name
    return {"id": f"C-2026-{no}", "kind": kind, "partner": name.split("_")[2], "title": "유지보수", "period": "2026-01-01 ~ 2026-12-31",
            "amount_text": "1,000,000원", "file": {"name": name, "url": FILE_BASE + name}}


# [기능명세] make_draft
# 목적: fixture PDF를 실제 extract로 읽어 문서 사전을 만들고 draft 계산
# 입력: (no: 번호 문자열, kind="수입")
# 출력: statement.draft 결과 사전
# 처리: compare.extract(PDF, "contract") → {"extraction": ...} → draft(문서, 계약 행)
# 예외/실패: PDF 판독 실패는 단언 실패
# 부수효과: PDF 읽기
def make_draft(no, kind="수입"):
    parsed = extract(fixture_pdf(no), "contract")
    assert parsed["error"] is None
    return draft({"extraction": parsed}, contract_row(no, kind))


# [기능명세] values_of
# 목적: 초안을 생성 요청 값으로 바꾸고 계약 정보(관리번호·계약기간)를 붙임
# 입력: (d: draft 결과, **over: 덮어쓸 값)
# 출력: build에 넘길 values 사전
# 처리: customer·issue_date·from/to·items 복사 + contract_no·period_start·period_end
# 예외/실패: 없음
# 부수효과: 없음
def values_of(d, **over):
    values = {"customer": d["customer"], "issue_date": d["issue_date"], "from_month": d["from_month"], "to_month": d["to_month"],
              "items": [dict(it) for it in d["items"]], "contract_no": d["contract_id"], "period_start": d["period_start"], "period_end": d["period_end"]}
    values.update(over)
    return values


# [기능명세] sha256
# 목적: 파일 내용 해시 계산(템플릿 원본 불변 확인용)
# 입력: (path: Path)
# 출력: 16진 문자열
# 처리: 파일 전체를 읽어 sha256
# 예외/실패: 파일 오류 전파
# 부수효과: 파일 읽기
def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


# [기능명세] test_draft_fixture_income_contracts
# 목적: 수입 계약 fixture 001·002·003의 초안 값(거래처·체결일·기간·품목·규격·금액·수입 판정)을 검증
# 입력: ()
# 출력: None; 조건을 어기면 AssertionError
# 처리: 세 계약서를 실제 extract로 읽어 draft 결과를 기대값과 비교
# 예외/실패: 단언 실패
# 부수효과: PDF 읽기
def test_draft_fixture_income_contracts():
    d = make_draft("002")
    assert d["eligible"] is True and d["reason"] is None
    assert d["customer"] == "세움바이오(주)" and d["customer_biz"] == "000-81-00012"
    assert d["contract_date"] == "2025-12-10" and d["issue_date"] == "2025-12-10"
    assert (d["period_start"], d["period_end"], d["from_month"], d["to_month"]) == ("2026-01-01", "2026-12-31", "2026-01", "2026-12")
    assert [(it["name"], it["spec"], it["monthly_amount"]) for it in d["items"]] == [("웹방화벽 서비스 유지보수", "인재 채용 홈페이지", 190000)]
    assert d["notes"] == [] and d["missing"] == [] and d["contract_id"] == "C-2026-002"

    d = make_draft("001")
    assert d["eligible"] and d["customer"] == "누리헬스케어(주)"
    assert (d["from_month"], d["to_month"]) == ("2026-01", "2026-12")
    assert d["items"][0]["monthly_amount"] == 580000 and d["items"][0]["spec"] == "대표 홈페이지와 고객 서비스 웹서버"

    d = make_draft("003")
    assert d["eligible"] and d["customer"] == "한결메디컬(주)" and d["contract_date"] == "2026-02-16"
    assert (d["from_month"], d["to_month"]) == ("2026-03", "2027-02")
    assert d["items"] == [{"name": "전자결재 시스템 운영 지원", "spec": "전자결재 서버 운영과 사용자 지원", "monthly_amount": 1250000}]


# [기능명세] test_draft_spend_contract_not_eligible
# 목적: 지출 계약 011(새봄바이오가 갑)은 eligible=false와 지출 안내 문구가 나오는지, 행 구분이 수입이라도 을 번호가 자사가 아니면 거절되는지 검증
# 입력: ()
# 출력: None; 조건을 어기면 AssertionError
# 처리: 011을 지출·수입 행으로 각각 draft
# 예외/실패: 단언 실패
# 부수효과: PDF 읽기
def test_draft_spend_contract_not_eligible():
    d = make_draft("011", "지출")
    assert d["eligible"] is False
    assert d["reason"] == "지출 계약은 거래처가 발행하는 문서라 거래명세서를 만들지 않습니다."
    assert make_draft("011", "수입")["eligible"] is False


# [기능명세] test_draft_multi_items_and_notes
# 목적: 이름이 붙은 서비스료 줄이 2개 이상이면 품목을 나누고 규격 누락·안내를 남기는지, 총액 계약·부가세 포함·체결일 없음 안내를 검증
# 입력: ()
# 출력: None; 조건을 어기면 AssertionError
# 처리: 최소 원문 텍스트와 fields로 만든 가짜 문서 두 개를 draft
# 예외/실패: 단언 실패
# 부수효과: 없음
def test_draft_multi_items_and_notes():
    text = ('"갑" : 가나(주)\n"을" : 주식회사 새봄바이오\n계약체결일 : [2026]년 [01]월 [05]일\n'
            "1. 서버 유지보수: 월 금 100,000원(부가가치세 별도)\n2. 백업 관리: 월 금 50,000원(부가가치세 별도)\n"
            "사업자등록번호 000-81-00099 000-81-00001\n")
    fields = {"start": {"value": "2026-01-01"}, "end": {"value": "2026-06-30"}, "basis": {"value": "monthly"},
              "amount": {"value": 100000}, "vat_mode": {"value": "excluded"}}
    d = draft({"extraction": {"fields": fields, "pages": [{"page": 1, "text": text}]}}, {"id": "C-1", "kind": "수입"})
    assert d["eligible"] and d["customer"] == "가나(주)" and d["contract_date"] == "2026-01-05"
    assert d["items"] == [{"name": "서버 유지보수", "spec": "", "monthly_amount": 100000}, {"name": "백업 관리", "spec": "", "monthly_amount": 50000}]
    assert d["missing"] == ["items[0].spec", "items[1].spec"]
    assert d["notes"] == ["서비스료 항목이 2개입니다. 각 품목의 규격을 견적서에서 확인해 입력하세요."]

    text = ('"갑" : 가나(주)\n별첨 1. 서비스의 범위\n1 서비스명 장비 납품\n2 대상 본사\n사업자등록번호 000-81-00099 000-81-00001\n')
    fields = {"basis": {"value": "total"}, "amount": {"value": 1100000}, "vat_mode": {"value": "included"}}
    d = draft({"extraction": {"fields": fields, "pages": [{"page": 1, "text": text}]}}, {"id": "C-2", "kind": "수입"})
    assert d["items"] == [{"name": "장비 납품", "spec": "본사", "monthly_amount": None}]
    assert d["notes"] == ["총액 계약입니다. 월금액을 직접 확인해 입력하세요.", "부가가치세 포함 금액입니다. 공급가액(부가세 제외)으로 고쳐 입력하세요."]
    assert d["contract_date"] is None and d["issue_date"] is None
    assert {"contract_date", "period_start", "period_end", "items[0].monthly_amount"} <= set(d["missing"])


# [기능명세] test_build_002_full_year
# 목적: 002 초안 그대로 만든 엑셀의 12줄 A~E 값, A3·A5 치환, F~H·25행 수식 보존, 계산 설정, 반환 총계, 템플릿 원본 불변을 검증
# 입력: (tmp_path)
# 출력: None; 조건을 어기면 AssertionError
# 처리: build → 결과 파일을 openpyxl로 다시 열어 칸 값 비교, 템플릿 해시 전후 비교
# 예외/실패: 단언 실패
# 부수효과: tmp_path에 xlsx 생성
def test_build_002_full_year(tmp_path):
    before = sha256(template_path())
    result = build(values_of(make_draft("002")), tmp_path / "out")
    assert result["file"] == "거래명세서_C-2026-002_세움바이오(주).xlsx"
    assert Path(result["path"]).is_file() and Path(result["path"]).parent == (tmp_path / "out").resolve()
    assert (result["rows"], result["supply"], result["vat"], result["total"], result["warnings"]) == (12, 2280000, 228000, 2508000, [])

    template = openpyxl.load_workbook(template_path()).active
    wb = openpyxl.load_workbook(result["path"])
    ws = wb.active
    assert wb.calculation.fullCalcOnLoad is True
    assert ws["A3"].value == "수    신  :  세움바이오(주)  貴 中"
    assert ws["A5"].value == "일      자  :  2025년 12월 10일"
    assert ws["A13"].value.date() == date(2026, 1, 31) and ws["A13"].number_format == "yyyy-mm-dd"
    assert (ws["B13"].value, ws["C13"].value, ws["D13"].value, ws["E13"].value) == ("웹방화벽 서비스 유지보수", "인재 채용 홈페이지", 190000, "1월")
    assert ws["A14"].value.date() == date(2026, 2, 28) and ws["E24"].value == "12월" and ws["A24"].value.date() == date(2026, 12, 31)
    assert ws["F13"].value == '=IF(D13="","",D13)'
    for r in range(13, 26):
        for col in "FGH":
            assert ws[f"{col}{r}"].value == template[f"{col}{r}"].value
    assert ws["F25"].value == "=SUM(F13:F24)" and ws["A25"].value == template["A25"].value
    for ref in ("E3", "F3", "F4", "F5", "F6", "F7", "F8", "H2"):
        assert ws[ref].value == template[ref].value
    assert sha256(template_path()) == before


# [기능명세] test_build_003_cross_year_labels
# 목적: 해가 바뀌는 003(2026-03~2027-02)의 청구월이 "26.03"~"27.02" 형식인지 검증
# 입력: (tmp_path)
# 출력: None; 조건을 어기면 AssertionError
# 처리: build 후 E13·E22·E24와 A24 날짜 확인
# 예외/실패: 단언 실패
# 부수효과: tmp_path에 xlsx 생성
def test_build_003_cross_year_labels(tmp_path):
    result = build(values_of(make_draft("003")), tmp_path)
    ws = openpyxl.load_workbook(result["path"]).active
    assert (ws["E13"].value, ws["E22"].value, ws["E24"].value) == ("26.03", "26.12", "27.02")
    assert ws["A24"].value.date() == date(2027, 2, 28) and ws["D13"].value == 1250000
    assert result["supply"] == 15000000 and result["file"] == "거래명세서_C-2026-003_한결메디컬(주).xlsx"


# [기능명세] test_build_rejects_without_saving
# 목적: 12줄 초과(품목 2개×7개월)·필수값 누락·from>to·0원은 ValueError이고 파일을 만들지 않는지 검증
# 입력: (tmp_path)
# 출력: None; 조건을 어기면 AssertionError
# 처리: 잘못된 값으로 build를 호출해 예외 문구와 폴더가 비어 있는지 확인
# 예외/실패: 단언 실패
# 부수효과: tmp_path 폴더 생성 가능(파일은 생성하지 않음)
def test_build_rejects_without_saving(tmp_path):
    base = make_draft("002")
    two = [{"name": "가", "spec": "A", "monthly_amount": 1}, {"name": "나", "spec": "B", "monthly_amount": 2}]
    with pytest.raises(ValueError, match="14줄"):
        build(values_of(base, items=two, from_month="2026-01", to_month="2026-07"), tmp_path)
    with pytest.raises(ValueError, match="규격"):
        build(values_of(base, items=[{"name": "가", "spec": " ", "monthly_amount": 1}]), tmp_path)
    with pytest.raises(ValueError, match="월금액"):
        build(values_of(base, items=[{"name": "가", "spec": "A", "monthly_amount": 0}]), tmp_path)
    with pytest.raises(ValueError, match="늦습니다"):
        build(values_of(base, from_month="2026-05", to_month="2026-04"), tmp_path)
    with pytest.raises(ValueError, match="수신"):
        build(values_of(base, customer=""), tmp_path)
    with pytest.raises(ValueError, match="품목을 1개"):
        build(values_of(base, items=[]), tmp_path)
    assert list(tmp_path.glob("*.xlsx")) == []
    # 2개×6개월 = 12줄은 허용된다.
    assert build(values_of(base, items=two, from_month="2026-01", to_month="2026-06"), tmp_path)["rows"] == 12


# [기능명세] test_build_same_name_and_partial_range
# 목적: 같은 값으로 두 번 만들면 두 번째가 _2로 저장되어 덮어쓰지 않는지, 일부 기간이면 범위 접미사와 경고가 붙는지 검증
# 입력: (tmp_path)
# 출력: None; 조건을 어기면 AssertionError
# 처리: 같은 값 build 두 번 → 파일명·두 파일 존재 확인 → 2026-01~06으로 build
# 예외/실패: 단언 실패
# 부수효과: tmp_path에 xlsx 3개 생성
def test_build_same_name_and_partial_range(tmp_path):
    values = values_of(make_draft("002"))
    first = build(values, tmp_path)
    second = build(values, tmp_path)
    assert first["file"] == "거래명세서_C-2026-002_세움바이오(주).xlsx"
    assert second["file"] == "거래명세서_C-2026-002_세움바이오(주)_2.xlsx"
    assert Path(first["path"]).is_file() and Path(second["path"]).is_file()
    part = build(values_of(make_draft("002"), from_month="2026-01", to_month="2026-06"), tmp_path)
    assert part["file"] == "거래명세서_C-2026-002_세움바이오(주)_202601-202606.xlsx"
    assert part["rows"] == 6 and len(part["warnings"]) == 1 and "계약기간" in part["warnings"][0]


# [기능명세] test_build_formula_injection_escaped
# 목적: =,+,-,@로 시작하는 품목·규격이 수식이 아니라 작은따옴표가 붙은 문자열로 저장되는지 검증
# 입력: (tmp_path)
# 출력: None; 조건을 어기면 AssertionError
# 처리: 품목 "=cmd", 규격 "@x"로 build 후 B13·C13 값과 데이터 형식 확인
# 예외/실패: 단언 실패
# 부수효과: tmp_path에 xlsx 생성
def test_build_formula_injection_escaped(tmp_path):
    result = build(values_of(make_draft("002"), items=[{"name": "=cmd", "spec": "@x", "monthly_amount": 10}]), tmp_path)
    ws = openpyxl.load_workbook(result["path"]).active
    assert ws["B13"].value == "'=cmd" and ws["C13"].value == "'@x" and ws["B13"].data_type == "s"


# [기능명세] unlock
# 목적: 실행 토큰과 비밀번호 인증을 마친 TestClient 준비
# 입력: (client: TestClient, token: 실행 토큰)
# 출력: None
# 처리: /?token= 으로 실행 쿠키 발급→config.APP_PASSWORD로 로그인
# 예외/실패: 인증 실패 시 AssertionError
# 부수효과: client 쿠키 설정
def unlock(client, token):
    assert client.get(f"/?token={token}").status_code == 200
    assert client.post("/api/auth/login", json={"password": config.APP_PASSWORD}).status_code == 200


# [기능명세] register
# 목적: fixture 계약서를 실제 판독 결과·실제 PDF 경로로 Store에 등록하고 가져온 계약 행 저장
# 입력: (store: Store, no: 번호 문자열, kind: "수입"|"지출")
# 출력: 저장한 계약 행
# 처리: extract → save_document(URL, 이름, 해시, 판독 결과, fixture 경로, kind="contract") → save_contract_import(행)
# 예외/실패: 파일·DB 오류 전파
# 부수효과: 임시 DB 갱신(fixture PDF는 읽기만 함)
def register(store, no, kind):
    item = contract_row(no, kind)
    pdf = fixture_pdf(no)
    store.save_document(item["file"]["url"], item["file"]["name"], sha256(pdf), extract(pdf, "contract"), pdf, kind="contract")
    store.save_contract_import(item)
    return item


# [기능명세] test_statement_api_flow
# 목적: 거래명세서 API의 초안 조회·생성·내려받기와 오류 응답(지출 400, 없는 계약 404, 원문 없음 400, 경로 조작 404, 작업 중 400)을 검증
# 입력: (tmp_path)
# 출력: None; 조건을 어기면 AssertionError
# 처리: Store(tmp_path/"data")에 002(수입)·011(지출)·문서 없는 계약 등록 → 잠금 해제 → GET 초안 → POST 생성 → 파일 존재·경로 확인
#       → GET /api/statements/{file} → 011 POST 400 → 없는 id 404 → 잘못된 파일명 404 → 배타 구간 중 POST 400
# 예외/실패: 단언 실패
# 부수효과: tmp_path 임시 DB·xlsx 생성(네트워크 없음)
def test_statement_api_flow(tmp_path):
    store = Store(tmp_path / "data")
    static = tmp_path / "static"
    static.mkdir()
    (static / "index.html").write_text("app")
    app = create_app(store, static, "tok")
    register(store, "002", "수입")
    register(store, "011", "지출")
    no_doc = contract_row("001", "수입")
    store.save_contract_import(no_doc)
    folder = (store.document_root() / "거래명세서").absolute()
    with TestClient(app, base_url=BASE) as client:
        assert client.get("/?token=tok").status_code == 200
        assert client.get("/api/contracts/C-2026-002/statement").status_code == 401
        unlock(client, "tok")

        r = client.get("/api/contracts/C-2026-002/statement")
        assert r.status_code == 200
        body = r.json()
        assert set(body) == {"contract_id", "eligible", "reason", "customer", "customer_biz", "contract_date", "issue_date", "period_start",
                             "period_end", "from_month", "to_month", "items", "notes", "missing", "folder"}
        assert body["eligible"] is True and body["customer"] == "세움바이오(주)" and body["folder"] == str(folder)
        assert body["items"] == [{"name": "웹방화벽 서비스 유지보수", "spec": "인재 채용 홈페이지", "monthly_amount": 190000}]

        request = {k: body[k] for k in ("customer", "issue_date", "from_month", "to_month", "items")}
        r = client.post("/api/contracts/C-2026-002/statement", json=request)
        assert r.status_code == 200, r.text
        result = r.json()
        assert set(result) == {"file", "path", "rows", "supply", "vat", "total", "warnings"}
        assert (result["rows"], result["supply"], result["vat"], result["total"]) == (12, 2280000, 228000, 2508000)
        assert Path(result["path"]).is_file() and Path(result["path"]).parent == folder.resolve()

        r = client.get(f"/api/statements/{result['file']}")
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert r.headers["content-disposition"].startswith("attachment") and r.content == Path(result["path"]).read_bytes()

        r = client.get("/api/contracts/C-2026-011/statement")
        assert r.status_code == 200 and r.json()["eligible"] is False and r.json()["reason"]
        r = client.post("/api/contracts/C-2026-011/statement", json=request)
        assert r.status_code == 400 and "지출 계약" in r.json()["detail"]
        r = client.post("/api/contracts/C-2026-002/statement", json={**request, "items": []})
        assert r.status_code == 400 and r.json()["detail"] == "품목을 1개 이상 입력해 주세요."

        assert client.get("/api/contracts/C-2099-999/statement").status_code == 404
        assert client.post("/api/contracts/C-2099-999/statement", json=request).status_code == 404
        r = client.get("/api/contracts/C-2026-001/statement")
        assert r.status_code == 400 and r.json()["detail"] == "계약서 원문 PDF가 없어 거래명세서를 만들 수 없습니다."

        (folder / "a.txt").write_text("x")
        for bad in ("../x.xlsx", "%2E%2E%2Fx.xlsx", "..%5Cx.xlsx", "a.txt", "없는파일.xlsx"):
            assert client.get(f"/api/statements/{bad}").status_code == 404, bad

        with app.state.service.exclusive("저장 위치 변경"):
            r = client.post("/api/contracts/C-2026-002/statement", json=request)
        assert r.status_code == 400 and "거래명세서 생성" in r.json()["detail"]
    assert len(list(folder.glob("*.xlsx"))) == 1
